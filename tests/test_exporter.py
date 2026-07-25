from __future__ import annotations

import datetime as dt
import io

import openpyxl
import pandas as pd

from app.exporter import export_to_csv_bytes, export_to_excel_bytes

SAMPLE_RECORDS = [
    {
        "id": 1,
        "group_name": "AI Jobs",
        "sender_name": "alice",
        "message_date": dt.datetime(2026, 7, 25, 10, 0, 0),
        "category": "Job",
        "title": "Backend Engineer",
        "summary": "Hiring a backend engineer.",
        "company": "Acme",
        "role": "Backend Engineer",
        "location": "Remote",
        "deadline": "2026-08-01",
        "apply_link": "https://acme.example/careers/123",
        "contact": None,
        "important_links": ["https://acme.example/careers/123"],
        "skills": ["Python", "SQL"],
        "tags": ["job"],
        "is_useful": True,
        "raw_text": "We are hiring a backend engineer",
    }
]


def test_export_to_csv_bytes_contains_header_and_row():
    content = export_to_csv_bytes(SAMPLE_RECORDS)
    text = content.decode("utf-8-sig")
    assert "title" in text.splitlines()[0]
    assert "Backend Engineer" in text
    assert "Python; SQL" in text


def test_export_to_excel_bytes_is_valid_workbook():
    content = export_to_excel_bytes(SAMPLE_RECORDS)
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook["Extractions"]
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert "title" in header
    values = [cell.value for cell in list(sheet.iter_rows())[1]]
    assert "Backend Engineer" in values


def test_export_to_csv_bytes_handles_empty_records():
    content = export_to_csv_bytes([])
    text = content.decode("utf-8-sig")
    assert "title" in text
    assert len(text.splitlines()) == 1
